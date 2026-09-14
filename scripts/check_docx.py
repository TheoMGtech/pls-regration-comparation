"""Structural gate for the Group 5 demand-management document."""

from __future__ import annotations

import sys
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document

REQUIRED_HEADERS = {
    "Data",
    "Integrante",
    "Demanda realizada",
    "Entrega ou evidencia",
    "Carga de trabalho",
    "Complexidade",
    "Status",
}


def normalized(value: str) -> str:
    return " ".join(value.strip().split()).replace("ê", "e").replace("é", "e")


def validate(path: Path) -> None:
    if not path.is_file():
        raise ValueError(f"document not found: {path}")
    try:
        with ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError("DOCX lacks word/document.xml")
    except BadZipFile as error:
        raise ValueError("file is not a valid DOCX package") from error

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "Grupo 5" not in text:
        raise ValueError("document must identify Grupo 5")
    headers = {
        normalized(cell.text)
        for table in document.tables
        for cell in table.rows[0].cells
    }
    missing = REQUIRED_HEADERS.difference(headers)
    if missing:
        raise ValueError(f"document lacks required demand headers: {sorted(missing)}")


if __name__ == "__main__":
    try:
        validate(Path(sys.argv[1] if len(sys.argv) == 2 else "docs/gestao-demandas-grupo-5.docx"))
    except (ValueError, KeyError, IndexError) as error:
        print(f"DOCX validation failed: {error}", file=sys.stderr)
        raise SystemExit(1)
    print("DOCX validation passed")
