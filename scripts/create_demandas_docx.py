"""Create the versioned Word template required for Group 5 demand tracking."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margin(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margin = properties.first_child_found_in("w:tcMar")
    if margin is None:
        margin = OxmlElement("w:tcMar")
        properties.append(margin)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margin.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margin.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_text(cell, text: str, bold: bool = False, size: int = 9, color: str = "000000") -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def build(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Gestão de Demandas do Grupo 5")
    title_run.font.name = "Aptos Display"
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    title_run.font.size = Pt(22)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.add_run("Finalidade. ").bold = True
    intro.add_run(
        "Registrar diariamente demandas executadas, responsabilidades e evidências do trabalho de "
        "Séries Temporais. Cada linha representa uma demanda concreta realizada em uma data."
    )

    info = document.add_table(rows=3, cols=2)
    info.style = "Table Grid"
    labels = [
        ("Grupo", "Grupo 5 - PLS Regression"),
        ("Integrantes e papéis", "Preencher antes do início das atividades"),
        ("Data de início", "Preencher"),
    ]
    for row, (label, value) in zip(info.rows, labels):
        row.cells[0].width = Cm(4.1)
        row.cells[1].width = Cm(12.0)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margin(cell)
        shade(row.cells[0], "D9EAF7")
        add_text(row.cells[0], label, bold=True)
        add_text(row.cells[1], value)

    document.add_paragraph()
    heading = document.add_paragraph()
    heading.add_run("Registro diário de demandas").bold = True
    heading.runs[0].font.size = Pt(13)

    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    headers = [
        "Data",
        "Integrante",
        "Demanda realizada",
        "Entrega ou evidência",
        "Carga de trabalho",
        "Complexidade",
        "Status",
    ]
    widths = [Cm(1.5), Cm(2.5), Cm(4.4), Cm(4.3), Cm(2.1), Cm(2.0), Cm(2.0)]
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(cell)
        shade(cell, "1F4E78")
        add_text(cell, header, bold=True, size=8, color="FFFFFF")
    set_repeat_table_header(table.rows[0])

    for _ in range(15):
        row = table.add_row()
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margin(cell, top=120, bottom=120)
            add_text(cell, "", size=8)

    document.add_paragraph()
    guidance = document.add_paragraph()
    guidance.add_run("Como preencher. ").bold = True
    guidance.add_run(
        "Descreva a atividade de forma objetiva e cite a evidência, como arquivo, seção, commit, "
        "resultado ou revisão. Carga mede esforço operacional; complexidade mede análise, novidade "
        "ou decisão técnica. Use Baixa, Média ou Alta para ambos; use A fazer, Em andamento ou Concluída no status."
    )
    warning = document.add_paragraph()
    warning.add_run("Não use registros genéricos como “ajudei no trabalho”.").bold = True

    document.core_properties.title = "Gestão de Demandas do Grupo 5"
    document.core_properties.subject = "Registro diário de demandas de Séries Temporais"
    document.core_properties.author = "Grupo 5"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


if __name__ == "__main__":
    build(Path("docs/gestao-demandas-grupo-5.docx"))
